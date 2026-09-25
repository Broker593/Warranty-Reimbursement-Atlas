#!/usr/bin/env python3
"""Build workpaper exports for the Warranty Atlas.

Outputs (under docs/downloads/):
  warranty-atlas.xlsx              one workbook: README, Summary, Coverage, Audit fields,
                                   Rate sample rules, Law dates, Weekly log, News
  state-pdfs/XX.pdf                one-page summary per state (citations + quotes)
  warranty-atlas-all-states.pdf    all 50 pages in one file
  state-pdfs.zip                   the 50 PDFs zipped

Inputs: docs/data/audit-fields.json, docs/research/labor-by-coverage-v3.json,
        docs/data/weekly-checks.json, docs/data/news.json
Requires: reportlab, openpyxl, pypdf   (pip install reportlab openpyxl pypdf)
Usage:    python3 tools/build_exports.py [--date YYYY-MM-DD]
"""
import argparse, datetime as dt, io, json, os, re, sys, zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(ROOT, 'docs')
OUT = os.path.join(DOCS, 'downloads')

APPLIES = {'yes': 'Required', 'conditional': 'Conditional', 'no': 'Not reached', 'silent': 'Not addressed', 'unverified': 'Unverified'}
HOURS = {'factory': 'Factory time (reasonable/adequate standard where stated)', 'multiplier': 'OEM time x multiplier',
         'independent_guide': 'Independent/retail time guide', 'actual_time': 'Actual technician time',
         'negotiated_other': 'Negotiated/other', 'silent': 'Statute silent on time / not established', 'n/a': '-'}
COV = [('factory_warranty', 'Factory warranty'), ('mfr_service_contract', 'Mfr-backed service contract'),
       ('cpo', 'CPO warranty'), ('independent_service_contract', 'Independent service contract')]
EXCL = {'MAINT': 'Routine maintenance', 'TIRES': 'Tires', 'ALIGN': 'Alignments', 'INSPECT': 'State inspections',
        'RECON': 'Prep / reconditioning', 'ACCESS': 'Accessories', 'BODY': 'Collision / body / glass', 'GOODWILL': 'Goodwill',
        'DISCOUNT': 'Discounted / promotional', 'FLEET': 'Fleet', 'GOVT': 'Government', 'INTERNAL': 'Internal',
        'SVC_CONTRACT': 'Service-contract paid', 'WARRANTY': 'Warranty / recall', 'INSURANCE': 'Insurance-paid',
        'NO_CHARGE': 'No-charge', 'ENGINE_TRANS': 'Engine / transmission assembly', 'DETAIL': 'Detailing', 'OTHER': 'Other (see text)'}
MODE = {'fewer': 'fewer of the two samples', 'greater_rate': 'sample with the higher rate', 'dealer_choice': 'dealer chooses',
        'ro_only': 'RO count only', 'calendar_month': 'prior calendar month', 'none': 'no statutory labor sample'}
# Enacted changes whose effective date is after the build date: show the pre-change value.
DATE_GATED = [
    {'state': 'RI', 'effective': '2026-10-01', 'cov': 'factory_warranty', 'field': 'paid_hours', 'before': 'silent',
     'label': 'Statute silent on time until 9/30/2026; actual technician time from 10/1/2026'},
    {'state': 'RI', 'effective': '2026-10-01', 'cov': 'cpo', 'field': 'applies', 'before': 'silent',
     'label': 'Not addressed until 9/30/2026; Required from 10/1/2026'},
]


def load(p):
    with open(os.path.join(DOCS, p), encoding='utf-8') as f:
        return json.load(f)


def trim(t, n):
    t = re.sub(r'\s+', ' ', str(t or '')).strip()
    return t if len(t) <= n else t[:n - 1].rstrip() + '…'


def yn(v, y='Yes', n='No'):
    return y if v is True else n if v is False else 'Silent'


def sample_text(a):
    c = a.get('calc') or {}
    if c.get('mode') == 'calendar_month':
        return 'All qualifying ROs, prior calendar month'
    if c.get('mode') == 'none':
        return 'No statutory labor-rate sample'
    bits = []
    if c.get('ro_count'): bits.append(f"{c['ro_count']} ROs")
    if c.get('days') and c.get('mode') != 'ro_only': bits.append(f"{c['days']} days")
    s = ' or '.join(bits) + f" ({MODE.get(c.get('mode'), '')})"
    if c.get('max_age_days'): s += f"; ROs <= {c['max_age_days']} days old"
    if c.get('denominator') == 'oem_hours': s += '; OEM-hour denominator'
    return s


def amended(a):
    d = a.get('law_dates') or {}
    if not d.get('last_amended_year'):
        return 'Unknown'
    s = str(d['last_amended_year'])
    if d.get('last_amendment_effective'): s += f" (eff. {d['last_amendment_effective']})"
    return s


def next_change(a):
    n = (a.get('law_dates') or {}).get('next_scheduled_change')
    return n if n and n.get('effective') else None


def coverage_cell(state, key, cov, build_date):
    applies, hours = cov['applies'], cov['paid_hours']
    note = ''
    for g in DATE_GATED:
        if g['state'] == state and g['cov'] == key and build_date < g['effective']:
            if g['field'] == 'applies': applies = g['before']
            else: hours = g['before']
            note = g['label']
    if applies in ('silent', 'no'):
        hours = 'n/a'
    return APPLIES.get(applies, applies), HOURS.get(hours, hours), note


# ---------------------------------------------------------------- Excel
def build_xlsx(audit, cov, weekly, news, build_date, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    F = Font(name='Arial', size=10); FB = Font(name='Arial', size=10, bold=True)
    FH = Font(name='Arial', size=10, bold=True, color='FFFFFF'); FT = Font(name='Arial', size=14, bold=True)
    fill = PatternFill('solid', fgColor='142943'); thin = Side(style='thin', color='9FB0C5')
    wrap = Alignment(wrap_text=True, vertical='top')

    def sheet(title, headers, rows, widths):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill, c.alignment = FH, fill, Alignment(wrap_text=True, vertical='center')
        for r in rows:
            ws.append(['' if v is None else v for v in r])
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font, c.alignment, c.border = F, wrap, Border(bottom=thin)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'C2' if len(headers) > 4 else 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 32
        return ws

    by = {r['state']: r for r in cov}
    ws = wb.active; ws.title = 'README'
    lines = [('Warranty Atlas · workpaper export', FT), (f'Built {build_date} from the live Atlas data files.', F), ('', F),
             ('What is in this workbook', FB),
             ('Summary: one row per state with coverage status, key audit fields, law dates and last-verified date. Count formulas at the right.', F),
             ('Coverage: 200 rows (50 states x 4 coverage types) with statute quotes and pinpoint cites.', F),
             ('Audit fields: claim deadlines, chargeback windows, rate submission, manufacturer response and challenge, penalties, with quotes.', F),
             ('Rate sample rules: each state\'s retail-rate sample, RO age limit and exclusions.', F),
             ('Law dates: last amendment, amending act, original enactment, next scheduled change.', F),
             ('Weekly log and News: the Weekly checks and News tabs.', F), ('', F),
             ('How to read it', FB),
             ('"Silent" / "Not addressed" means the reviewed statute says nothing. Contracts, regulations or other law may still apply.', F),
             ('"Conditional" coverage usually turns on who the legal obligor is, or whether the manufacturer pays for the work. Program branding does not establish the obligor.', F),
             ('Rhode Island: enacted changes take effect 10/1/2026. Cells show the rule in force on the build date and name the upcoming change.', F), ('', F),
             ('Limits', FB),
             ('Public-source research summary for audit planning. It is not legal advice and does not contain SOA-approved rates or dealer payment data. Confirm against the cited statute before relying on a specific rule; escalate disputes to Legal.', F),
             ('Sources: official legislature/revisor text unless the record says enrolled act or mirror. Quotes were machine-checked against saved source text during research.', F)]
    for i, (t, f) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t); c.font = f; c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 130

    srows = []
    for a in audit:
        s = a['state']; c = by.get(s, {}).get('coverage', {})
        cells = {k: coverage_cell(s, k, c[k], build_date) for k, _ in COV if k in c}
        n = next_change(a)
        srows.append([s, a['name'], cells['factory_warranty'][1], cells['mfr_service_contract'][0], cells['cpo'][0],
                      cells['independent_service_contract'][0], a['claims'].get('decision_deadline_days'),
                      yn(a['claims'].get('deemed_approved_if_late')), a['claims'].get('payment_deadline_days'),
                      a['chargebacks'].get('lookback_months'), a['rate_submission'].get('frequency_limit') or 'Silent',
                      sample_text(a), a['manufacturer_response'].get('response_deadline_days'),
                      yn(a['manufacturer_response'].get('deemed_approved_if_no_response')), amended(a),
                      (n['effective'] + ' ' + (n.get('act') or '')) if n else '', a['verified']['audit_fields'], a.get('confidence')])
    ws = sheet('Summary', ['State', 'Name', 'Factory paid hours', 'Mfr service contract', 'CPO', 'Independent SC',
                           'Claim decision (days)', 'Late claim deemed approved', 'Payment (days)', 'Chargeback window (months)',
                           'Rate resubmission', 'Rate sample', 'Mfr response (days)', 'No response = approved',
                           'Law last amended', 'Next scheduled change', 'Last verified', 'Confidence'],
               srows, [7, 15, 30, 16, 14, 16, 11, 12, 10, 12, 22, 34, 11, 12, 18, 26, 12, 11])
    # count block (formulas) to the right of the table
    last = len(srows) + 1
    ws['T1'] = 'Counts'; ws['T1'].font = FB
    block = [('Mfr SC Required', f'=COUNTIF($D$2:$D${last},"Required")'), ('Mfr SC Conditional', f'=COUNTIF($D$2:$D${last},"Conditional")'),
             ('CPO Required', f'=COUNTIF($E$2:$E${last},"Required")'), ('CPO Conditional', f'=COUNTIF($E$2:$E${last},"Conditional")'),
             ('Indep SC Conditional', f'=COUNTIF($F$2:$F${last},"Conditional")'),
             ('Late claims deemed approved', f'=COUNTIF($H$2:$H${last},"Yes")'),
             ('Chargeback window 12 months', f'=COUNTIF($J$2:$J${last},12)'),
             ('Chargeback window under 12 months', f'=COUNTIFS($J$2:$J${last},"<12",$J$2:$J${last},">0")'),
             ('No response = approved', f'=COUNTIF($N$2:$N${last},"Yes")'),
             ('States with a scheduled change', f'=COUNTIF($P$2:$P${last},"?*")')]
    for i, (lab, fml) in enumerate(block, 2):
        ws.cell(row=i, column=20, value=lab).font = F
        ws.cell(row=i, column=21, value=fml).font = F
    ws.column_dimensions['T'].width = 32; ws.column_dimensions['U'].width = 8

    crow = []
    for r in cov:
        for k, lab in COV:
            x = r['coverage'][k]; st, hrs, note = coverage_cell(r['state'], k, x, build_date)
            crow.append([r['state'], lab, st, hrs, note, r.get('hourly_rate_method') if k == 'factory_warranty' else '',
                         x.get('quote'), x.get('pinpoint'), r.get('primary_cite'), r.get('confidence'), r.get('source_quality'), r.get('official_url')])
    sheet('Coverage', ['State', 'Coverage type', 'Status', 'Paid hours', 'Date note', 'Hourly rate method (factory)', 'Statute quote',
                       'Pinpoint', 'Primary cite', 'Confidence', 'Source type', 'Official URL'],
          crow, [7, 24, 14, 30, 26, 50, 60, 22, 28, 11, 12, 40])

    arow = []
    for a in audit:
        c, cb, rs, mr, pe = a['claims'], a['chargebacks'], a['rate_submission'], a['manufacturer_response'], a['penalties']
        arow.append([a['state'], a['name'], c.get('decision_deadline_days'), yn(c.get('deemed_approved_if_late')), c.get('payment_deadline_days'),
                     c.get('payment_deadline_trigger'), c.get('dealer_filing_deadline'), c.get('resubmission_rights'), c.get('denial_requirements'),
                     c.get('quote'), c.get('pinpoint'), cb.get('lookback_months'), cb.get('fraud_extension'), cb.get('limits'), cb.get('quote'), cb.get('pinpoint'),
                     rs.get('frequency_limit'), sample_text(a), rs.get('formula'), '; '.join(EXCL.get(x, x) for x in rs.get('exclusions') or []),
                     rs.get('exclusions_text'), rs.get('new_rate_effective'), rs.get('quote'), rs.get('pinpoint'),
                     mr.get('response_deadline_days'), yn(mr.get('deemed_approved_if_no_response')), mr.get('challenge_standard'), mr.get('challenge_method'),
                     mr.get('dispute_forum'), mr.get('quote'), mr.get('pinpoint'), pe.get('private_remedies'), pe.get('admin_sanctions'), pe.get('cite'),
                     '; '.join(a.get('cites') or []), a.get('official_url'), a.get('source_quality'), a.get('confidence'), a['verified']['audit_fields'], a.get('notes')])
    sheet('Audit fields', ['State', 'Name', 'Claim decision (days)', 'Late claim deemed approved', 'Payment (days)', 'Payment trigger', 'Dealer filing deadline',
                           'Resubmission rights', 'Denial requirements', 'Claims quote', 'Claims pinpoint', 'Chargeback window (months)', 'Fraud extension',
                           'Chargeback limits', 'Chargeback quote', 'Chargeback pinpoint', 'Rate resubmission frequency', 'Rate sample', 'Rate formula',
                           'Excluded from sample', 'Exclusions (statute wording)', 'New rate effective', 'Rate quote', 'Rate pinpoint',
                           'Mfr response (days)', 'No response = approved', 'Challenge standard', 'Challenge method', 'Dispute forum', 'Response quote',
                           'Response pinpoint', 'Private remedies', 'Admin sanctions', 'Penalty cites', 'All cites', 'Official URL', 'Source type',
                           'Confidence', 'Last verified', 'Notes'],
          arow, [7, 15] + [16] * 38)

    rrow = []
    for a in audit:
        c = a['calc']
        rrow.append([a['state'], a['name'], MODE.get(c['mode'], c['mode']), c.get('ro_count'), c.get('days'), c.get('max_age_days'),
                     'OEM time-allowance hours' if c.get('denominator') == 'oem_hours' else 'Hours billed',
                     '; '.join(c.get('exclusions') or []), a['rate_submission'].get('exclusions_text'), a['rate_submission'].get('formula'),
                     a['rate_submission'].get('frequency_limit'), c.get('note')])
    sheet('Rate sample rules', ['State', 'Name', 'Sample rule', 'RO count', 'Day window', 'Max RO age (days)', 'Denominator', 'Exclusion codes',
                                'Exclusions (statute wording)', 'Formula', 'Resubmission frequency', 'Note'],
          rrow, [7, 15, 24, 9, 10, 12, 20, 30, 50, 40, 22, 40])

    lrow = []
    for a in audit:
        d = a.get('law_dates') or {}; n = next_change(a) or {}
        lrow.append([a['state'], a['name'], d.get('section'), d.get('last_amended_year'), d.get('last_amending_act'), d.get('last_amendment_effective'),
                     d.get('originally_enacted_year'), n.get('effective'), n.get('act'), n.get('summary'), a['verified']['audit_fields'],
                     a['verified']['coverage'], d.get('history_source')])
    sheet('Law dates', ['State', 'Name', 'Section', 'Last amended (year)', 'Amending act', 'Amendment effective', 'Originally enacted', 'Next change effective',
                        'Next change act', 'Next change summary', 'Audit fields verified', 'Coverage verified', 'History source'],
          lrow, [7, 15, 26, 12, 34, 14, 12, 14, 30, 50, 13, 13, 60])

    wrow = []
    for e in sorted(weekly.get('entries', []), key=lambda x: x['check_date'], reverse=True):
        for it in e.get('items') or [{}]:
            wrow.append([e['check_date'], e.get('headline'), it.get('category'), it.get('state'), it.get('bill'), it.get('effective'),
                         it.get('status'), it.get('summary_short') or it.get('summary'), it.get('atlas_impact'), it.get('source_url')])
    sheet('Weekly log', ['Check date', 'Headline', 'Category', 'State', 'Bill / act', 'Effective', 'Status', 'Summary', 'Atlas impact', 'Source'],
          wrow, [12, 40, 16, 7, 30, 12, 40, 60, 40, 40])

    nrow = []
    for w in sorted(news.get('weeks', []), key=lambda x: x['week_of'], reverse=True):
        for n in w.get('items', []):
            nrow.append([w['week_of'], n.get('date'), n.get('title'), n.get('publisher'), n.get('source_type'), n.get('perspective'),
                         ', '.join(n.get('states') or []), n.get('why_it_matters'), n.get('url')])
    sheet('News', ['Week of', 'Date', 'Title', 'Publisher', 'Source type', 'Perspective', 'States', 'Why it matters', 'URL'],
          nrow, [12, 12, 50, 30, 16, 16, 10, 60, 40])
    wb.save(path)


# ---------------------------------------------------------------- PDFs
def build_pdf(a, c, build_date, path, scale=1.0):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape as X
    navy = colors.HexColor('#142943'); grey = colors.HexColor('#9FB0C5')
    fs = 7.4 * scale
    B = ParagraphStyle('b', fontName='Helvetica', fontSize=fs, leading=fs * 1.22, alignment=TA_LEFT)
    BB = ParagraphStyle('bb', parent=B, fontName='Helvetica-Bold')
    H = ParagraphStyle('h', parent=B, fontName='Helvetica-Bold', fontSize=fs + 2.2, leading=(fs + 2.2) * 1.2, textColor=navy, spaceBefore=4, spaceAfter=2)
    T = ParagraphStyle('t', parent=B, fontName='Helvetica-Bold', fontSize=15, leading=18, textColor=colors.white)
    TS = ParagraphStyle('ts', parent=B, fontSize=7.8, leading=10, textColor=colors.white)
    Q = ParagraphStyle('q', parent=B, fontName='Helvetica-Oblique', leftIndent=6, textColor=colors.HexColor('#1f2d40'))
    P = lambda t, s=B: Paragraph(X(str(t if t is not None else '')), s)
    L = int(260 * scale)

    n = next_change(a)
    story = []
    head = Table([[Paragraph(f"WARRANTY ATLAS · {X(a['name'])} ({a['state']})", T)],
                  [Paragraph(X(f"Law last amended: {amended(a)}  ·  Next scheduled change: {(n['effective'] + ' — ' + trim(n.get('act'), 60)) if n else 'none found'}  ·  "
                               f"Last verified: {a['verified']['audit_fields']}  ·  Built {build_date}"), TS)]],
                 colWidths=[7.6 * inch])
    head.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), navy), ('LEFTPADDING', (0, 0), (-1, -1), 8), ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5)]))
    story += [head, Spacer(1, 4), P('Cites: ' + '; '.join(a.get('cites') or []), B)]

    def grid(rows, widths, header=True):
        t = Table(rows, colWidths=[w * inch for w in widths], repeatRows=1 if header else 0)
        st = [('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LINEBELOW', (0, 0), (-1, -1), 0.4, grey), ('LEFTPADDING', (0, 0), (-1, -1), 3),
              ('RIGHTPADDING', (0, 0), (-1, -1), 3), ('TOPPADDING', (0, 0), (-1, -1), 1.6), ('BOTTOMPADDING', (0, 0), (-1, -1), 1.6)]
        if header: st += [('BACKGROUND', (0, 0), (-1, 0), navy), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white)]
        t.setStyle(TableStyle(st)); return t
    HW = ParagraphStyle('hw', parent=BB, textColor=colors.white)

    story.append(Paragraph('Labor reimbursement by coverage type', H))
    rows = [[P('Coverage', HW), P('Status', HW), P('Paid hours', HW), P('Pinpoint / date note', HW)]]
    for k, lab in COV:
        x = c['coverage'][k]; st, hrs, note = coverage_cell(a['state'], k, x, build_date)
        rows.append([P(lab, BB), P(st, BB), P(hrs), P(trim((x.get('pinpoint') or '') + ('  ·  ' + note if note else ''), 110))])
    story.append(grid(rows, [1.6, 1.0, 2.3, 2.7]))
    story.append(P('Hourly rate: ' + trim(c.get('hourly_rate_method'), int(480 * scale)), B))

    cl, cb, rs, mr, pe = a['claims'], a['chargebacks'], a['rate_submission'], a['manufacturer_response'], a['penalties']
    d = lambda v: f'{v} days' if v is not None else 'Silent'
    story.append(Paragraph('Audit fields', H))
    rows = [[P('Field', HW), P('Rule', HW), P('Pinpoint', HW)],
            [P('Claim decision', BB), P(f"{d(cl.get('decision_deadline_days'))}; late claim deemed approved: {'Yes' if cl.get('deemed_approved_if_late') is True else 'no rule in statute'}. " + trim(cl.get("denial_requirements"), L)), P(trim(cl.get('pinpoint'), 60))],
            [P('Payment', BB), P(f"{d(cl.get('payment_deadline_days'))}{' ' + cl['payment_deadline_trigger'] if cl.get('payment_deadline_trigger') else ''}. Dealer filing deadline: {cl.get('dealer_filing_deadline') or 'Silent'}. " + trim(cl.get('resubmission_rights'), L // 2)), P('')],
            [P('Chargebacks', BB), P(f"{str(cb['lookback_months']) + ' months' if cb.get('lookback_months') is not None else 'Silent'}; fraud: {trim(cb.get('fraud_extension') or 'Silent', 70)}. " + trim(cb.get('limits'), L)), P(trim(cb.get('pinpoint'), 60))],
            [P('Rate resubmission', BB), P(f"{rs.get('frequency_limit') or 'Silent'}. Sample: {sample_text(a)}. Excluded: {', '.join(EXCL.get(x, x) for x in rs.get('exclusions') or []) or 'none listed'}."), P(trim(rs.get('pinpoint'), 60))],
            [P('Mfr response', BB), P(f"{d(mr.get('response_deadline_days'))}; no response = approved: {'Yes' if mr.get('deemed_approved_if_no_response') is True else 'no rule in statute'}. Challenge: " + trim(mr.get('challenge_standard'), L) + (f" Forum: {trim(mr.get('dispute_forum'), 70)}." if mr.get('dispute_forum') else '')), P(trim(mr.get('pinpoint'), 60))],
            [P('Penalties', BB), P(trim((pe.get('private_remedies') or '') + ' ' + (pe.get('admin_sanctions') or ''), int(L * 1.3))), P(trim(pe.get('cite'), 70))]]
    story.append(grid(rows, [1.15, 5.15, 1.3]))

    story.append(Paragraph('Statute quotes', H))
    qs = [('Claims', cl), ('Chargebacks', cb), ('Rate', rs), ('Factory time', c['coverage']['factory_warranty'])]
    for lab, blk in qs:
        if blk.get('quote'):
            story.append(Paragraph(f"<b>{lab}</b> ({X(trim(blk.get('pinpoint'), 60))}): “{X(trim(blk['quote'], int(420 * scale)))}”", Q))
    if a.get('notes'):
        story.append(Paragraph('Notes', H)); story.append(P(trim(a['notes'], int(900 * scale))))
    story.append(Spacer(1, 4))
    story.append(P(f"Source: {a.get('source_quality')} · {trim(a.get('official_url'), 120)} · Confidence: {a.get('confidence')}. "
                   "Public-source research summary for audit planning, not legal advice. No SOA-approved rates. Confirm against the cited statute; escalate disputes to Legal.",
                   ParagraphStyle('f', parent=B, fontSize=fs - 0.8, leading=(fs - 0.8) * 1.2, textColor=colors.HexColor('#3d4c60'))))
    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.45 * inch, rightMargin=0.45 * inch, topMargin=0.4 * inch, bottomMargin=0.4 * inch,
                            title=f"Warranty Atlas - {a['name']}", author='Warranty Atlas')
    doc.build(story)


def pages(path):
    from pypdf import PdfReader
    return len(PdfReader(path).pages)


def main():
    ap = argparse.ArgumentParser(); ap.add_argument('--date', default=dt.date.today().isoformat())
    build_date = ap.parse_args().date
    audit = load('data/audit-fields.json'); cov = load('research/labor-by-coverage-v3.json')
    weekly = load('data/weekly-checks.json'); news = load('data/news.json')
    by = {r['state']: r for r in cov}
    missing = sorted(set(a['state'] for a in audit) ^ set(by))
    if missing: sys.exit(f'State mismatch between audit and coverage data: {missing}')
    os.makedirs(os.path.join(OUT, 'state-pdfs'), exist_ok=True)
    build_xlsx(audit, cov, weekly, news, build_date, os.path.join(OUT, 'warranty-atlas.xlsx'))
    from pypdf import PdfReader, PdfWriter
    allw = PdfWriter(); zbuf = io.BytesIO(); z = zipfile.ZipFile(zbuf, 'w', zipfile.ZIP_DEFLATED)
    for a in audit:
        p = os.path.join(OUT, 'state-pdfs', a['state'] + '.pdf')
        for scale in (1.0, 0.93, 0.86, 0.8, 0.74):
            build_pdf(a, by[a['state']], build_date, p, scale)
            if pages(p) == 1: break
        else:
            sys.exit(f"{a['state']}: could not fit on one page")
        allw.add_page(PdfReader(p).pages[0]); z.write(p, a['state'] + '.pdf')
    with open(os.path.join(OUT, 'warranty-atlas-all-states.pdf'), 'wb') as f:
        allw.write(f)
    z.close()
    with open(os.path.join(OUT, 'state-pdfs.zip'), 'wb') as f:
        f.write(zbuf.getvalue())
    print(f'Built exports for {len(audit)} states ({build_date}).')


if __name__ == '__main__':
    main()
